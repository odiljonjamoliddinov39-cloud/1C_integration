"""Excel export of any table and bulk invoice entry from an uploaded workbook."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Company, Counterparty, Invoice, Item
from app.services import invoices as invoice_service

BULK_COLUMNS = ["invoice_no", "date", "buyer_inn", "contract_number", "item", "quantity", "price"]


def to_xlsx(title: str, columns: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for idx, col in enumerate(columns, start=1):
        width = max([len(str(col))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r)])
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(width + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def bulk_template() -> bytes:
    sample = [
        ["1", date.today().isoformat(), "123456789", "15", "00000000000000000", 10, 15000],
        ["1", date.today().isoformat(), "123456789", "15", "Вода питьевая 19л", 5, ""],
    ]
    return to_xlsx("invoices", BULK_COLUMNS, sample)


def _cell_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def parse_bulk(db: Session, company: Company, content: bytes) -> list[dict]:
    """Group rows into invoices and validate them without saving anything.

    Returns [{"key", "date", "buyer_ref", "buyer_name", "buyer_inn", "contract_ref", "rows", "total", "vat",
    "errors": [...]}]; `errors` also carries problems resolving buyers, contracts and items.
    """
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h or "").strip().lower() for h in rows[0]]
    missing = [c for c in BULK_COLUMNS if c not in header and c != "price"]
    if missing:
        raise invoice_service.InvoiceError(f"Missing columns: {', '.join(missing)}")
    col = {name: header.index(name) for name in BULK_COLUMNS if name in header}

    cps = list(db.scalars(select(Counterparty).where(Counterparty.company_id == company.id, Counterparty.deleted.is_(False))))
    by_inn: dict[str, Counterparty] = {}
    for cp in cps:
        by_inn.setdefault(cp.inn, cp)
    items = list(db.scalars(select(Item).where(Item.company_id == company.id, Item.deleted.is_(False))))
    by_ikpu = {i.ikpu_code: i for i in items if i.ikpu_code}
    by_name = {i.name.strip().lower(): i for i in items}

    groups: dict[str, dict] = {}
    for line_no, raw in enumerate(rows[1:], start=2):
        if not raw or all(v in (None, "") for v in raw):
            continue

        def get(name):
            return raw[col[name]] if name in col and col[name] < len(raw) else None

        key = str(get("invoice_no") or "").strip()
        group = groups.setdefault(
            key,
            {"key": key, "lines": [], "errors": [], "date": None, "buyer_inn": "", "contract_number": ""},
        )
        group["lines"].append(line_no)
        when = _cell_date(get("date"))
        if when is None:
            group["errors"].append({"line": line_no, "field": "date", "message": "Date is missing or invalid"})
        group["date"] = group["date"] or when
        group["buyer_inn"] = group["buyer_inn"] or str(get("buyer_inn") or "").strip()
        group["contract_number"] = group["contract_number"] or str(get("contract_number") or "").strip()

        item_key = str(get("item") or "").strip()
        item = by_ikpu.get(item_key) or by_name.get(item_key.lower())
        if item is None:
            group["errors"].append({"line": line_no, "field": "item", "message": f"Item «{item_key}» not found"})
            continue
        price = get("price")
        group.setdefault("rows", []).append(
            {"item_ref": item.ref_1c, "quantity": str(get("quantity") or 0), "price": "" if price in (None, "") else str(price)}
        )

    result = []
    for key, group in groups.items():
        buyer = by_inn.get(group["buyer_inn"])
        contract_ref = None
        if buyer is None:
            group["errors"].append({"field": "buyer_inn", "message": f"No counterparty with INN {group['buyer_inn']}"})
        else:
            contracts = buyer.contract_refs or []
            if group["contract_number"]:
                match = [c for c in contracts if str(c.get("number", "")).strip() == group["contract_number"]]
                contract_ref = match[0]["ref"] if match else None
            elif len(contracts) == 1:
                contract_ref = contracts[0]["ref"]
        draft = Invoice(
            company_id=company.id,
            date=group["date"] or date.today(),
            buyer_ref=buyer.ref_1c if buyer else None,
            buyer_inn=group["buyer_inn"],
            buyer_name=buyer.name if buyer else "",
            contract_ref=contract_ref,
            rows=invoice_service.fill_rows_from_items(db, company.id, group.get("rows", [])),
        )
        invoice_service.recalc(draft)
        errors = group["errors"] + invoice_service.validate(db, draft)
        result.append(
            {
                "key": key,
                "lines": group["lines"],
                "date": draft.date.isoformat(),
                "buyer_ref": draft.buyer_ref,
                "buyer_inn": draft.buyer_inn,
                "buyer_name": draft.buyer_name,
                "contract_ref": draft.contract_ref,
                "rows": draft.rows,
                "total": str(draft.total),
                "vat": str(draft.vat),
                "errors": errors,
            }
        )
    return result
